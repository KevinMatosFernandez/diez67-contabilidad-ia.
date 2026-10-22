
from pathlib import Path
from openpyxl import load_workbook


COLUMNAS_REQUERIDAS = [
    "Producto",
    "Inicio",
    "Entrada",
    "Rotura",
    "consumo",
    "Disponible",
    "Final",
    "Precio",
    "Precio sug 30% CUP",
    "Precio/usd",
    "Precio sug 30%",
    "Venta",
    "Importe",
    "Importe/usd",
    "Costo",
    "Costo/usd",
    "Importe Costo",
    "Importe Costo/usd",
    "Ganancias",
    "Ganancias/usd",
    "Margen (%)",
    "costo rotura",
    "costo consumo",
    "Tasa Ref",
]


def leer_excel(ruta_archivo):
    ruta = Path(ruta_archivo)

    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {ruta}")

    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = libro.active

    encabezados = [
        str(celda.value).strip() if celda.value is not None else ""
        for celda in next(hoja.iter_rows(min_row=1, max_row=1))
    ]

    faltantes = [
        nombre for nombre in COLUMNAS_REQUERIDAS
        if nombre not in encabezados
    ]

    if faltantes:
        libro.close()
        raise ValueError(
            "Faltan estas columnas en el Excel: "
            + ", ".join(faltantes)
        )

    indice_producto = encabezados.index("Producto")
    productos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if indice_producto >= len(fila):
            continue

        nombre = fila[indice_producto]

        if nombre is not None and str(nombre).strip():
            productos.append(str(nombre).strip())

    libro.close()

    if len(productos) != len(set(productos)):
        raise ValueError(
            "Hay nombres de productos repetidos en el archivo."
        )

    return productos


def main():
    print("DIEZ67 MARKET — Importador de Excel")
    ruta = input("Introduce la ruta del archivo .xlsx: ").strip().strip('"')

    try:
        productos = leer_excel(ruta)
        print(f"Excel leído correctamente. Productos encontrados: {len(productos)}")
        for producto in productos[:10]:
            print(f"- {producto}")

        if len(productos) > 10:
            print(f"... y {len(productos) - 10} productos más.")

    except Exception as error:
        print(f"No se pudo leer el Excel: {error}")


if __name__ == "__main__":
    main()

